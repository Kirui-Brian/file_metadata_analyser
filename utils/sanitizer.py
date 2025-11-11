"""
Metadata Sanitizer Utility
Removes or anonymizes metadata from files for privacy protection.
"""

import os
import shutil
from pathlib import Path
from typing import Optional
from datetime import datetime

try:
    from PIL import Image
    PIL_AVAILABLE = True
except ImportError:
    PIL_AVAILABLE = False

try:
    import PyPDF2
    import fitz  # PyMuPDF
    PDF_AVAILABLE = True
except ImportError:
    PDF_AVAILABLE = False

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

try:
    from openpyxl import load_workbook
    XLSX_AVAILABLE = True
except ImportError:
    XLSX_AVAILABLE = False

try:
    from pptx import Presentation
    PPTX_AVAILABLE = True
except ImportError:
    PPTX_AVAILABLE = False


class MetadataSanitizer:
    """
    Removes or sanitizes metadata from various file types.
    """
    
    def __init__(self):
        """Initialize the sanitizer."""
        pass
    
    def sanitize_file(self, input_path: str, output_path: Optional[str] = None) -> bool:
        """
        Remove metadata from a file based on its type.
        
        Args:
            input_path (str): Path to input file
            output_path (str, optional): Path for output file. If None, overwrites original.
            
        Returns:
            bool: True if successful, False otherwise
        """
        input_file = Path(input_path)
        
        if not input_file.exists():
            print(f"Error: File not found: {input_path}")
            return False
        
        if output_path is None:
            # Create backup
            backup_path = str(input_file) + '.backup'
            shutil.copy2(input_path, backup_path)
            output_path = input_path
            print(f"Backup created: {backup_path}")
        
        ext = input_file.suffix.lower()
        
        try:
            if ext in ['.jpg', '.jpeg', '.png', '.tiff', '.tif']:
                return self._sanitize_image(input_path, output_path)
            elif ext == '.pdf':
                return self._sanitize_pdf(input_path, output_path)
            elif ext == '.docx':
                return self._sanitize_docx(input_path, output_path)
            elif ext == '.xlsx':
                return self._sanitize_xlsx(input_path, output_path)
            elif ext == '.pptx':
                return self._sanitize_pptx(input_path, output_path)
            else:
                print(f"Warning: Sanitization not supported for {ext} files")
                # Just copy the file
                if input_path != output_path:
                    shutil.copy2(input_path, output_path)
                return True
        except Exception as e:
            print(f"Error sanitizing file: {str(e)}")
            return False
    
    def _sanitize_image(self, input_path: str, output_path: str) -> bool:
        """Remove EXIF data from images."""
        if not PIL_AVAILABLE:
            print("Error: PIL/Pillow not available")
            return False
        
        try:
            # Open image
            image = Image.open(input_path)
            
            # Remove EXIF data by saving without it
            # Create new image without metadata
            data = list(image.getdata())
            image_without_exif = Image.new(image.mode, image.size)
            image_without_exif.putdata(data)
            
            # Save without metadata
            save_params = {}
            if image.format == 'JPEG':
                save_params['quality'] = 95
                save_params['optimize'] = True
            
            image_without_exif.save(output_path, **save_params)
            
            print(f"✓ Image metadata removed: {output_path}")
            return True
        
        except Exception as e:
            print(f"Error removing image metadata: {str(e)}")
            return False
    
    def _sanitize_pdf(self, input_path: str, output_path: str) -> bool:
        """Remove metadata from PDF files."""
        if not PDF_AVAILABLE:
            print("Error: PDF libraries not available")
            return False
        
        try:
            # Use PyMuPDF (fitz) for better metadata handling
            doc = fitz.open(input_path)
            
            # Clear metadata
            doc.set_metadata({
                'title': '',
                'author': '',
                'subject': '',
                'keywords': '',
                'creator': '',
                'producer': '',
                'creationDate': '',
                'modDate': ''
            })
            
            # Save cleaned PDF
            doc.save(output_path, garbage=4, deflate=True, clean=True)
            doc.close()
            
            print(f"✓ PDF metadata removed: {output_path}")
            return True
        
        except Exception as e:
            print(f"Error removing PDF metadata: {str(e)}")
            # Try fallback method with PyPDF2
            try:
                return self._sanitize_pdf_pypdf2(input_path, output_path)
            except:
                return False
    
    def _sanitize_pdf_pypdf2(self, input_path: str, output_path: str) -> bool:
        """Fallback PDF sanitization using PyPDF2."""
        try:
            with open(input_path, 'rb') as input_file:
                pdf_reader = PyPDF2.PdfReader(input_file)
                pdf_writer = PyPDF2.PdfWriter()
                
                # Copy all pages
                for page in pdf_reader.pages:
                    pdf_writer.add_page(page)
                
                # Remove metadata by not adding it
                pdf_writer.add_metadata({})
                
                # Write to output
                with open(output_path, 'wb') as output_file:
                    pdf_writer.write(output_file)
            
            print(f"✓ PDF metadata removed (PyPDF2): {output_path}")
            return True
        
        except Exception as e:
            print(f"Error with PyPDF2 method: {str(e)}")
            return False
    
    def _sanitize_docx(self, input_path: str, output_path: str) -> bool:
        """Remove metadata from DOCX files."""
        if not DOCX_AVAILABLE:
            print("Error: python-docx not available")
            return False
        
        try:
            doc = Document(input_path)
            
            # Clear core properties
            core_props = doc.core_properties
            core_props.author = ''
            core_props.title = ''
            core_props.subject = ''
            core_props.keywords = ''
            core_props.comments = ''
            core_props.last_modified_by = ''
            core_props.category = ''
            core_props.created = datetime.now()
            core_props.modified = datetime.now()
            
            # Save cleaned document
            doc.save(output_path)
            
            print(f"✓ DOCX metadata removed: {output_path}")
            return True
        
        except Exception as e:
            print(f"Error removing DOCX metadata: {str(e)}")
            return False
    
    def _sanitize_xlsx(self, input_path: str, output_path: str) -> bool:
        """Remove metadata from XLSX files."""
        if not XLSX_AVAILABLE:
            print("Error: openpyxl not available")
            return False
        
        try:
            # Load workbook
            wb = load_workbook(input_path)
            
            # Clear core properties
            wb.properties.creator = ''
            wb.properties.lastModifiedBy = ''
            wb.properties.title = ''
            wb.properties.subject = ''
            wb.properties.description = ''
            wb.properties.keywords = ''
            wb.properties.category = ''
            wb.properties.comments = ''
            wb.properties.created = datetime.now()
            wb.properties.modified = datetime.now()
            
            # Save cleaned workbook
            wb.save(output_path)
            wb.close()
            
            print(f"✓ XLSX metadata removed: {output_path}")
            return True
        
        except Exception as e:
            print(f"Error removing XLSX metadata: {str(e)}")
            return False
    
    def _sanitize_pptx(self, input_path: str, output_path: str) -> bool:
        """Remove metadata from PPTX files."""
        if not PPTX_AVAILABLE:
            print("Error: python-pptx not available")
            return False
        
        try:
            # Load presentation
            prs = Presentation(input_path)
            
            # Clear core properties
            core_props = prs.core_properties
            core_props.author = ''
            core_props.title = ''
            core_props.subject = ''
            core_props.keywords = ''
            core_props.comments = ''
            core_props.last_modified_by = ''
            core_props.category = ''
            core_props.created = datetime.now()
            core_props.modified = datetime.now()
            
            # Save cleaned presentation
            prs.save(output_path)
            
            print(f"✓ PPTX metadata removed: {output_path}")
            return True
        
        except Exception as e:
            print(f"Error removing PPTX metadata: {str(e)}")
            return False
    
    def sanitize_directory(self, directory: str, output_directory: Optional[str] = None) -> dict:
        """
        Sanitize all supported files in a directory.
        
        Args:
            directory (str): Input directory
            output_directory (str, optional): Output directory
            
        Returns:
            dict: Results summary
        """
        results = {
            'total': 0,
            'success': 0,
            'failed': 0,
            'skipped': 0
        }
        
        input_dir = Path(directory)
        if not input_dir.exists():
            print(f"Error: Directory not found: {directory}")
            return results
        
        # Create output directory if specified
        if output_directory:
            output_dir = Path(output_directory)
            output_dir.mkdir(parents=True, exist_ok=True)
        else:
            output_dir = input_dir
        
        # Process all files
        for file_path in input_dir.rglob('*'):
            if not file_path.is_file():
                continue
            
            results['total'] += 1
            
            # Determine output path
            relative_path = file_path.relative_to(input_dir)
            if output_directory:
                out_path = output_dir / relative_path
                out_path.parent.mkdir(parents=True, exist_ok=True)
            else:
                out_path = file_path
            
            # Sanitize file
            if self.sanitize_file(str(file_path), str(out_path)):
                results['success'] += 1
            else:
                results['failed'] += 1
        
        return results
    
    @staticmethod
    def create_sanitization_report(input_path: str, output_path: str) -> str:
        """
        Create a report comparing file sizes before and after sanitization.
        
        Args:
            input_path (str): Original file path
            output_path (str): Sanitized file path
            
        Returns:
            str: Report text
        """
        input_size = Path(input_path).stat().st_size
        output_size = Path(output_path).stat().st_size
        size_diff = input_size - output_size
        size_diff_pct = (size_diff / input_size * 100) if input_size > 0 else 0
        
        report = f"""
Sanitization Report
{'='*50}
Original File:  {input_path}
Original Size:  {input_size:,} bytes
Cleaned File:   {output_path}
Cleaned Size:   {output_size:,} bytes
Size Reduction: {size_diff:,} bytes ({size_diff_pct:.2f}%)
{'='*50}
"""
        return report
